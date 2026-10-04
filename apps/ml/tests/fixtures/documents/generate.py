# SPDX-License-Identifier: AGPL-3.0-only
"""Writes the sample documents used by the parser tests (all text is our own).

Usage, from apps/ml: uv run python tests/fixtures/documents/generate.py
"""

from pathlib import Path
from typing import cast

from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.shapes.autoshape import Shape

HERE = Path(__file__).parent
TITLE = "Surefy Sample Handbook"
HEADING = "Getting started"
PARAGRAPH = "Every workspace keeps its documents in knowledge bases that chat can cite."
SCANNED_TEXT = "Scanned invoice number 4711"
TABLE = [["Plan", "Seats"], ["Team", "10"], ["Business", "50"]]


def _pdf(pages: list[list[tuple[int, int, int, str]]]) -> bytes:
    """A minimal PDF with Helvetica text; each line is (size, x, y, text)."""
    objects: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"",  # page tree, filled in below
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    kids: list[int] = []
    for lines in pages:
        stream = b"".join(
            f"BT /F1 {size} Tf {x} {y} Td ({text}) Tj ET\n".encode() for size, x, y, text in lines
        )
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"endstream")
        content = len(objects)
        objects.append(
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
            b"/Resources << /Font << /F1 3 0 R >> >> /Contents %d 0 R >>" % content
        )
        kids.append(len(objects))
    refs = " ".join(f"{kid} 0 R" for kid in kids)
    objects[1] = f"<< /Type /Pages /Kids [{refs}] /Count {len(kids)} >>".encode()

    out = bytearray(b"%PDF-1.7\n")
    offsets: list[int] = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % offset for offset in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1,
        xref,
    )
    return bytes(out)


def text_pdf() -> None:
    page_one = [(24, 72, 760, TITLE), (16, 72, 710, HEADING), (11, 72, 680, PARAGRAPH)]
    page_two = [(16, 72, 760, "Limits"), (11, 72, 730, "Each organization may upload files.")]
    (HERE / "text.pdf").write_bytes(_pdf([page_one, page_two]))


def scanned_pdf() -> None:
    image = Image.new("RGB", (1240, 1754), "white")
    font = ImageFont.load_default(size=56)
    ImageDraw.Draw(image).text((120, 200), SCANNED_TEXT, fill="black", font=font)
    image.save(HERE / "scanned.pdf", resolution=150)


def docx() -> None:
    document = Document()
    document.add_heading(TITLE, level=0)
    document.add_heading(HEADING, level=1)
    document.add_paragraph(PARAGRAPH)
    table = document.add_table(rows=len(TABLE), cols=len(TABLE[0]))
    for row, values in zip(table.rows, TABLE, strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    document.save(str(HERE / "sample.docx"))


def xlsx() -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Plans"
    for values in TABLE:
        sheet.append(values)
    workbook.save(HERE / "sample.xlsx")


def pptx() -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    title, body = (cast("Shape", slide.placeholders[index]) for index in (0, 1))
    title.text_frame.text = HEADING
    body.text_frame.text = PARAGRAPH
    presentation.save(str(HERE / "sample.pptx"))


def html() -> None:
    rows = "".join(
        "<tr>" + "".join(f"<td>{value}</td>" for value in values) + "</tr>" for values in TABLE
    )
    (HERE / "sample.html").write_text(
        f"<!doctype html><html><head><title>{TITLE}</title></head><body>"
        f"<h1>{TITLE}</h1><h2>{HEADING}</h2><p>{PARAGRAPH}</p>"
        f"<table>{rows}</table></body></html>\n"
    )


if __name__ == "__main__":
    for write in (text_pdf, scanned_pdf, docx, xlsx, pptx, html):
        write()
